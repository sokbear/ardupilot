# -*- coding: utf-8 -*-
"""Конвертация Markdown → Word (.docx) с заголовками, таблицами и списками."""
from __future__ import annotations

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

from bench_diagrams import cache_path_for_mermaid, render_mermaid_to_png


def set_doc_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_diagram_image(doc: Document, mermaid_source: str, cache_dir: str) -> None:
    png_path = cache_path_for_mermaid(mermaid_source, cache_dir)
    if not render_mermaid_to_png(mermaid_source, png_path):
        doc.add_paragraph("【Блок-схема: не удалось построить графику】", style="Intense Quote")
        add_code_block(doc, mermaid_source.splitlines())
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(png_path, width=Inches(6.5))
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_sep(line: str) -> bool:
    s = line.strip().replace("|", "").replace(":", "").replace("-", "").strip()
    return len(s) == 0 and "|" in line


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.rows[ri].cells[ci]
            text = row[ci] if ci < len(row) else ""
            cell.text = text
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(10)
                    if ri == 0:
                        run.bold = True
    doc.add_paragraph()


def md_to_docx(md_path: str, docx_path: str) -> None:
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    here = os.path.dirname(os.path.abspath(md_path))
    diagram_cache = os.path.join(here, ".diagram_cache")
    os.makedirs(diagram_cache, exist_ok=True)

    doc = Document()
    set_doc_defaults(doc)

    title = os.path.splitext(os.path.basename(md_path))[0].replace("_", " ")
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    i = 0
    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    list_mode: str | None = None  # 'ul' or 'ol'

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                if code_lang == "mermaid":
                    add_diagram_image(doc, "\n".join(code_lines), diagram_cache)
                else:
                    add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
                code_lang = ""
            else:
                in_code = True
                code_lang = stripped[3:].strip().lower()
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            list_mode = None
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            table_rows = [parse_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_rows.append(parse_table_row(lines[i].strip()))
                i += 1
            add_table(doc, table_rows)
            list_mode = None
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2), level=level)
            list_mode = None
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("─" * 40)
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"`([^`]+)`", r"\1", text)
            doc.add_paragraph(text, style="List Bullet")
            list_mode = "ul"
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            doc.add_paragraph(text, style="List Number")
            list_mode = "ol"
            i += 1
            continue

        text = stripped
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)

        p = doc.add_paragraph()
        parts = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = p.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            else:
                p.add_run(part)
        list_mode = None
        i += 1

    doc.save(docx_path)
    print(f"OK: {docx_path}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="*", help="Markdown files")
    parser.add_argument("-o", "--out-dir", default=None, help="Output directory")
    args = parser.parse_args()

    here = os.path.dirname(os.path.abspath(__file__))
    inputs = args.inputs or [
        os.path.join(here, "SOFTWARE_DESCRIPTION.md"),
        os.path.join(here, "TRACKING_ALGORITHM_SPEC.md"),
        os.path.join(here, "README.md"),
    ]
    out_dir = args.out_dir or here

    for md in inputs:
        if not os.path.isfile(md):
            print(f"SKIP (not found): {md}", file=sys.stderr)
            continue
        base = os.path.splitext(os.path.basename(md))[0]
        docx = os.path.join(out_dir, base + ".docx")
        md_to_docx(md, docx)
    return 0


if __name__ == "__main__":
    sys.exit(main())
